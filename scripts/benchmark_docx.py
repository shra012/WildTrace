#!/usr/bin/env python3
"""Word documents for the setup comparison (run r2).

    python scripts/benchmark_docx.py spec      # docs/benchmarks/eval_suite_r2.docx (metric definitions)
    python scripts/benchmark_docx.py results   # docs/benchmarks/results_r2.docx (from summary.json)

Metric text comes from wildtrace/eval_metrics.METRICS, so the document and
the code that computes the numbers share one definition.
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from benchmark_models import REPO_ROOT, provenance  # noqa: E402
from benchmark_setups import comparison, evalset_path, out_dir  # noqa: E402

from wildtrace import eval_metrics as em  # noqa: E402

DIRECTION = {"higher": "↑ higher", "lower": "↓ lower", "target": "◎ target", "info": "info"}
LEVEL = {"attempt": "per drawing", "subject": "per subject", "setup": "per setup"}


def new_document(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    doc.add_heading(title, level=0)
    meta = doc.add_paragraph(subtitle)
    meta.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return doc


def _shade(cell: Any, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def table(doc: Document, header: list[str], rows: list[list[Any]], widths_cm: list[float], font_pt: float = 9) -> None:
    grid = doc.add_table(rows=1, cols=len(header))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for cell, text, width in zip(grid.rows[0].cells, header, widths_cm):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(font_pt)
        cell.width = Cm(width)
        _shade(cell, "DCE6F0")
    for values in rows:
        cells = grid.add_row().cells
        for cell, value, width in zip(cells, values, widths_cm):
            cell.text = ""
            cell.paragraphs[0].add_run("" if value is None else str(value)).font.size = Pt(font_pt)
            cell.width = Cm(width)
    doc.add_paragraph()


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


# ── spec ───────────────────────────────────────────────────────────────────

def build_spec() -> Path:
    cfg = comparison()
    evalset = json.loads(evalset_path().read_text())
    subjects = evalset["subjects"]
    prov = provenance()
    model = em.calibrate_draw_time(REPO_ROOT / "isaac-sim")
    doc = new_document(
        "WildTrace eval suite: setup comparison r2",
        f"{date.today().isoformat()} · git {prov['git_sha']} · config {prov['config_hash']} · "
        "code: wildtrace/eval_metrics.py, scripts/benchmark_setups.py",
    )

    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(
        "WildTrace turns animal photos into simple line drawings, then into pen trajectories that an xArm7 "
        "draws in Isaac Sim. Each drawing comes out of a loop: a drawing model generates it, an OpenCV "
        "prescreen and a vision-language judge check it, and a failed drawing is retried with adjusted "
        "settings. This suite compares two complete setups (drawing model plus judge) on the same subjects, "
        "with the same measurements, so we can decide which setup produces better training data, at what "
        "cost, and whether the judges can be trusted."
    )

    doc.add_heading("2. Setups compared", level=1)
    rows = []
    for setup in cfg["setups"]:
        validator = setup["validator"]
        judge = validator.get("openrouter_model_name") or validator.get("ollama_model_name")
        where = "OpenRouter API" if "openrouter" in validator["backend"] else "local Ollama"
        overrides = ", ".join(f"{k}={v}" for k, v in setup.get("rectifier_overrides", {}).items()) or "production settings"
        rows.append([setup["name"], setup["rectifier"], overrides, f"{judge} ({where})"])
    table(doc, ["Setup", "Drawing model", "Drawing settings", "Judge"], rows, [3.6, 2.6, 5.2, 5.6])
    bullets(doc, [
        "omnigen2 is a unified model: a Qwen2.5-VL-3B encoder reads the instruction and a 4B diffusion decoder "
        "redraws the masked subject crop as a line drawing. It runs in bf16 with model CPU offload. The crop is "
        "letterboxed onto a white square and the drawing cropped back, because on elongated crops OmniGen2 "
        "returned blank pages. The instruction is an edit command ('Convert this photo into a black and white line "
        "drawing ... keep the same pose, size and position'), and a judge's retry prompt is appended as a style note.",
        "flux-schnell is FLUX.1-schnell (GGUF Q4) img2img: it starts from the subject's silhouette contour and "
        "refines it in 4 steps. It was the production drawing model before this comparison. In its smoke run, "
        "keeping FLUX loaded while Ollama loaded the judge was OOM-killed (24.4 GB RAM, 6.3 GB swap), so this "
        "setup unloads FLUX after every generation. That changes its latency (a reload per attempt) but not its drawings.",
        f"Both setups run the production loop with at most {cfg['max_attempts']} attempts per subject.",
    ])

    doc.add_heading("3. Sample set", level=1)
    counts = Counter(s["category"] for s in subjects)
    buckets = Counter(s["angle_bucket"] for s in subjects)
    doc.add_paragraph(
        f"{len(subjects)} subjects drawn with seed {evalset['seed']} from the silver subject set: "
        f"{cfg['evalset']['per_category']} per category, plus "
        + ", ".join(f"{n} extra {c}" for c, n in cfg["evalset"].get("extra", {}).items())
        + ". Both setups draw exactly these subjects."
    )
    table(doc, ["Category", "Subjects"], [[c, n] for c, n in sorted(counts.items())], [5, 3])
    table(doc, ["Angle bucket", "Subjects"], [[b, n] for b, n in sorted(buckets.items())], [5, 3])
    if evalset.get("excluded_categories"):
        doc.add_paragraph(
            f"Excluded: {', '.join(evalset['excluded_categories'])}. The configuration lists it, but OpenImages "
            "v5 has no segmentation masks for Butterfly (or for any insect class), and the fetch stage requires "
            "masks, so no Butterfly subjects exist. Butterfly is still one of the answer options in the BioCLIP "
            "category check, so it acts as a distractor there."
        )

    doc.add_heading("4. Protocol", level=1)
    bullets(doc, [
        "Setups run one after another, never at the same time, so their memory use can't interfere. "
        "A watcher samples system RAM, swap and GPU once a second, and an out-of-memory kill is recorded as a result.",
        "Recording wrappers log every generation, judge call and retry-feedback call to attempts.ndjson "
        "(parameters, the prompt actually used, latency, verdict, tokens, cost). The production loop code is not modified.",
        "After both runs: every attempt's image is measured (section 5), and every final drawing is re-judged "
        "by the same three referees, whichever setup produced it.",
        "BioCLIP and the referees run only after generation has finished, so they never share memory with the drawing models.",
    ])
    doc.add_heading("Fairness notes", level=2)
    bullets(doc, [
        "Each setup's own judge decides its accept/retry loop, so acceptance rates reflect judge strictness as well "
        "as drawing quality. The cross-referee pass rates (group F) are the like-for-like quality comparison.",
        "Retry prompts are not treated the same way. Retry feedback from the judge suggests a new prompt and "
        "strength. OmniGen2 uses the suggested prompt; FLUX-schnell always rebuilds its fixed prompt, so its retries only change "
        "strength. Both are measured as production runs them, and attempts.ndjson records the prompt each attempt used.",
        "A first OmniGen2 run was stopped after 11 subjects and discarded: the silver stage's isolated image is "
        "the full frame, not the crop, so OmniGen2 got a misaligned, partly erased subject, and its first attempt "
        "was blank on every subject. Both bugs were fixed before the reported run (see section 2).",
        "A second OmniGen2 run used one fixed seed for every attempt. OmniGen2 returns a blank page for some "
        "seeds, so a subject that came out blank stayed blank on every retry, whereas FLUX draws fresh noise on each "
        "call. The reported run seeds each attempt from its subject and attempt number, which is reproducible. "
        "The fixed-seed run is kept in outputs/benchmarks/r2/archive_A_fixed_seed (7/20 accepted).",
        "Inputs differ by design. OmniGen2 reads the masked photo, while FLUX starts from the silhouette contour. "
        "That gives FLUX a head start on silhouette fidelity (group C); that group should be read alongside "
        "recognizability (group E).",
    ])

    doc.add_heading("5. Metrics", level=1)
    doc.add_paragraph(
        "Direction: ↑ higher is better, ↓ lower is better, ◎ closer to a target is better, info = descriptive. "
        "Per-drawing metrics are reported on each setup's final drawing (and, where noted, on all attempts)."
    )
    for group, title in em.GROUPS.items():
        doc.add_heading(f"{group}. {title}", level=2)
        rows = [[m.name, DIRECTION[m.direction], LEVEL[m.level], m.formula, m.meaning]
                for m in em.METRICS if m.group == group]
        table(doc, ["Metric", "Better", "Level", "Definition", "What it tells you"], rows, [3.2, 1.7, 1.8, 5.2, 5.1], font_pt=8.5)
    doc.add_heading("How the less obvious metrics work", level=2)
    bullets(doc, [
        "Filled silhouette: the drawing's outer contours are closed with a 5 px kernel and filled, which turns an "
        "outline into a solid shape that can be compared with the source photo's mask.",
        "Source mask: the subject's OpenImages mask, cropped with the same box as the photo and resized to the drawing.",
        "Pose: the pipeline's own viewpoint classifier (wildtrace/viewpoint.py) is run on the filled drawing, the same "
        "way it originally labelled the photo. A mismatch means the gold record's angle label would be wrong for the drawing.",
        "Trajectory: the exact gold-stage export (_build_trajectory_payload: up to 6 outer contours, 16–128 points "
        "each) is scaled onto the 160×120 mm paper the way the simulator does it (one uniform scale).",
    ])
    if model:
        bullets(doc, [
            f"Draw-time model: duration ≈ {model.a_s_per_mm:.3f} s/mm × pen-down path + {model.b_s_per_stroke:.1f} s/stroke "
            f"+ {model.c_s:.1f} s. Fitted on {model.n} Isaac Sim runs made with the current drawing configuration "
            f"(R² = {model.r2}). It is an estimate for comparing setups, not a measured robot time.",
        ])

    doc.add_heading("6. Statistics", level=1)
    bullets(doc, [
        "Rates and means come with 95% bootstrap confidence intervals (2,000 resamples of subjects, seed 14).",
        "Paired comparison: for each subject, the setup with the better final drawing wins on a given metric; "
        "an exact two-sided sign test on wins versus losses (ties dropped) gives the p-value.",
        "Judge agreement: Cohen's κ between each pair of referees on pass/fail over the same drawings "
        "(κ < 0.2 slight, 0.2–0.4 fair, 0.4–0.6 moderate, > 0.6 substantial).",
        "With 20 subjects, a difference in rates smaller than about 20 percentage points is usually within noise. "
        "Treat such results as directional only.",
    ])

    doc.add_heading("7. Limitations", level=1)
    bullets(doc, [
        "No human labels: 'quality' here means automated metrics plus model judges. The strong referee "
        "(Qwen3-VL-235B) is a proxy, not ground truth.",
        "Six categories, 20 subjects; Butterfly is not covered.",
        "Draw time is estimated from a 7-run fit, not simulated for each drawing.",
        "OmniGen2 is run with tuned settings (image guidance 2.2, 20 steps, guidance on 60% of steps); "
        "FLUX is run with its production settings.",
    ])

    path = REPO_ROOT / "docs" / "benchmarks" / f"eval_suite_{cfg['run_id']}.docx"
    doc.save(path)
    return path


# ── results ────────────────────────────────────────────────────────────────

def _fmt(cell: Any) -> str:
    if isinstance(cell, dict):
        if cell.get("value") is None and "p50" in cell:
            return f"p50 {cell['p50']} / p95 {cell['p95']}"
        value, ci = cell.get("value"), cell.get("ci")
        if value is None:
            return "–"
        return f"{value}  [{ci[0]}–{ci[1]}]" if ci and ci[0] is not None else str(value)
    return "–" if cell is None else str(cell)


def build_results() -> Path:
    cfg = comparison()
    summary = json.loads((out_dir() / "summary.json").read_text())
    setups = list(summary["setups"])
    registry = {m.key: m for m in em.METRICS}
    doc = new_document(
        f"WildTrace setup comparison {summary['run_id']}: results",
        f"{date.today().isoformat()} · git {summary['git_sha']} · config {summary['config_hash']} · "
        f"definitions: eval_suite_{summary['run_id']}.docx",
    )
    width = [5.0] + [12.0 / len(setups)] * len(setups)

    findings = REPO_ROOT / "docs" / "benchmarks" / f"{summary['run_id']}_findings.md"
    if findings.exists():
        doc.add_heading("Key findings", level=1)
        for line in findings.read_text().splitlines():
            if not line.startswith("- "):
                continue
            paragraph = doc.add_paragraph(style="List Bullet")
            for i, part in enumerate(line[2:].split("**")):
                paragraph.add_run(part).bold = i % 2 == 1

    def metric_rows(keys: list[str], source: str | None) -> list[list[str]]:
        rows = []
        for key in keys:
            name = registry[key].name if key in registry else key
            arrow = {"higher": " ↑", "lower": " ↓", "target": " ◎"}.get(registry[key].direction, "") if key in registry else ""
            rows.append([name + arrow] + [_fmt(summary["setups"][s][source][key] if source else summary["setups"][s].get(key))
                                          for s in setups])
        return rows

    doc.add_heading("1. Outcome, judges and resources", level=1)
    keys = ["accepted", "first_attempt_pass", "attempts", "combined_score", "error", "subject_latency_s",
            "own_judge_pass", "parse_failed", "cost_usd", "peak_vram_gb", "sys_ram_gb"]
    rows = metric_rows(keys, None)
    rows += [["Generation time per attempt (s)"] + [_fmt(summary["setups"][s]["gen_latency_s"]) for s in setups],
             ["Judge latency (s)"] + [_fmt(summary["setups"][s]["judge_latency_s"]) for s in setups],
             ["Swap used (GB)"] + [_fmt(summary["setups"][s].get("swap_gb")) for s in setups],
             ["OOM kills (auto-resumed)"] + [_fmt(summary["setups"][s].get("oom_kills")) for s in setups],
             ["Wall time (s)"] + [_fmt(summary["setups"][s].get("wall_s")) for s in setups]]
    table(doc, ["Metric"] + setups, rows, width)

    doc.add_heading("2. Cross-referee pass rate (same judges for both setups)", level=1)
    referees = list(summary["setups"][setups[0]]["referee"])
    table(doc, ["Referee"] + setups,
          [[r + (" (strong)" if any(x.get("strong") and x["name"] == r for x in cfg["referees"]) else "")]
           + [_fmt(summary["setups"][s]["referee"][r]) for s in setups] for r in referees], width)
    table(doc, ["Referee pair", "Cohen's κ", "Raw agreement", "n"],
          [[k, v["kappa"], v["agreement"], v["n"]] for k, v in summary["referee_agreement"].items()], [7, 3, 3, 2])

    doc.add_heading("3. Final-drawing metrics", level=1)
    for group, title in em.GROUPS.items():
        keys = [k for k in summary["setups"][setups[0]]["final"] if registry[k].group == group]
        if keys:
            doc.add_heading(f"{group}. {title}", level=2)
            table(doc, ["Metric"] + setups, metric_rows(keys, "final"), width)

    doc.add_heading("4. Paired per-subject comparison", level=1)
    table(doc, ["Metric", "Wins: " + setups[0], "Wins: " + setups[1], "Ties", "Sign-test p"],
          [[k, v["a_wins"], v["b_wins"], v["ties"], v["sign_test_p"]] for k, v in summary["paired"].items()],
          [5, 3.2, 3.2, 2, 2.6])

    doc.add_heading("5. Per category", level=1)
    for setup in setups:
        doc.add_heading(setup, level=2)
        table(doc, ["Category", "n", "Accepted", "Silhouette IoU", "BioCLIP top-1", "Strong-referee pass"],
              [[c, v["n"], v["accepted"], v["silhouette_iou"], v["clip_top1"], v["strong_referee_pass"]]
               for c, v in summary["per_category"][setup].items()], [3, 1.2, 2.4, 3, 3, 3.4])

    sheet = out_dir() / "contact_sheet.png"
    if sheet.exists():
        doc.add_heading("6. Contact sheet", level=1)
        doc.add_paragraph("Source photo, then each setup's final drawing. Label: accepted/rejected, attempts, "
                          "silhouette IoU, strong-referee score.")
        # One tall image would run off the page; add it 5 subjects at a time.
        from PIL import Image

        image = Image.open(sheet)
        header, row, per_page = 24, 180, 5
        for top in range(header, image.height, row * per_page):
            chunk = out_dir() / f"contact_sheet_{top}.png"
            image.crop((0, top, image.width, min(image.height, top + row * per_page))).save(chunk)
            doc.add_picture(str(chunk), width=Cm(11))

    path = REPO_ROOT / "docs" / "benchmarks" / f"results_{summary['run_id']}.docx"
    doc.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("document", choices=["spec", "results"])
    args = parser.parse_args()
    print(build_spec() if args.document == "spec" else build_results())


if __name__ == "__main__":
    main()

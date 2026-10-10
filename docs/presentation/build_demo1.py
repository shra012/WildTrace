"""Build the WildTrace Demo 1 slide deck.

Numbers are copied from docs/workbook/build_workbook.py, docs/benchmarks/r2.md,
and isaac-sim/WildTrace_RL_Isaac_Sim_Overview.md. Do not invent new measurements.
"""
from __future__ import annotations

import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FIG = Path(__file__).resolve().parent / "figures"
OUT = Path(__file__).resolve().parent / "WildTrace_Demo1.pptx"

NAVY = RGBColor(0x0E, 0x2A, 0x47)
BLUE = RGBColor(0x1F, 0x6F, 0xC4)
ORANGE = RGBColor(0xC4, 0x4E, 0x1A)
INK = RGBColor(0x1A, 0x1D, 0x21)
MUTED = RGBColor(0x5C, 0x65, 0x70)
CARD = RGBColor(0xF4, 0xF7, 0xFB)
LINE = RGBColor(0xD7, 0xDE, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xE7, 0xF0, 0xFA)

BLUE_HEX, ORANGE_HEX = "#1F6FC4", "#E07A3D"
INK_HEX, MUTED_HEX, GRID_HEX = "#1A1D21", "#5C6570", "#E6E8EC"

W, H = Inches(13.333), Inches(7.5)
SLIDES = 14


def _font(run, size, bold=False, color=INK, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._r.get_or_add_rPr().set(qn("a:latin"), name) if False else None
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = rpr.makeelement(qn("a:latin"), {})
        rpr.append(latin)
    latin.set("typeface", name)


def _box(slide, l, t, w, h, text, size=16, bold=False, color=INK, align="left", anchor="top"):
    shape = slide.shapes.add_textbox(l, t, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    _font(run, size, bold, color)
    return shape


def _lines(slide, l, t, w, h, rows, size=15, color=INK, bold=False, spacing=6):
    shape = slide.shapes.add_textbox(l, t, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, row in enumerate(rows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0)
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = row
        _font(run, size, bold, color)
    return shape


def _rect(slide, l, t, w, h, fill, line=None, radius=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    # Keep cards behind text added later.
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)
    return shape


def _footer(slide, page):
    _box(slide, Inches(0.5), Inches(7.18), Inches(10.2), Inches(0.24),
         "WildTrace  ·  Demo 1  ·  figures from Workbook 1, benchmark r2, and baseline-v1",
         11, color=MUTED)
    _box(slide, Inches(11.4), Inches(7.18), Inches(1.45), Inches(0.24),
         f"{page}  /  {SLIDES}", 11, color=MUTED, align="right")


def _header(slide, kicker, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.08), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    _box(slide, Inches(0.48), Inches(0.22), Inches(12.2), Inches(0.28), kicker.upper(), 12, True, ORANGE)
    _box(slide, Inches(0.48), Inches(0.46), Inches(12.3), Inches(0.48), title, 28, True, NAVY)


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def _style_cell(cell, text, size, bold, color, fill, align="left"):
    cell.text = ""
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT if align == "left" else PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    _font(run, size, bold, color)
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill


def _table(slide, l, t, w, h, headers, rows, col_widths, size=13, header_fill=NAVY):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), l, t, w, h)
    table = table_shape.table
    for i, width in enumerate(col_widths):
        table.columns[i].width = width
    for j, header in enumerate(headers):
        _style_cell(table.cell(0, j), header, size, True, WHITE, header_fill)
    for r, row in enumerate(rows):
        fill = WHITE if r % 2 == 0 else CARD
        for c, value in enumerate(row):
            _style_cell(table.cell(r + 1, c), value, size, c == 0, INK, fill)
    return table_shape


def _kpi(slide, l, t, w, h, value, label):
    _rect(slide, l, t, w, h, CARD, radius=0.12)
    _box(slide, l + Inches(0.14), t + Inches(0.1), w - Inches(0.22), Inches(0.46),
         value, 26, True, BLUE)
    _box(slide, l + Inches(0.14), t + Inches(0.58), w - Inches(0.22), Inches(0.48),
         label, 13, False, MUTED)


# ── figures ──────────────────────────────────────────────────────────────

def _mpl():
    plt.rcParams.update({
        "font.family": "Calibri",
        "font.size": 13,
        "axes.edgecolor": "#C5CDD6",
        "axes.labelcolor": MUTED_HEX,
        "xtick.color": MUTED_HEX,
        "ytick.color": INK_HEX,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.dpi": 200,
        "savefig.bbox": "tight",
        "savefig.pad_inches": 0.15,
    })


def make_figures() -> dict[str, Path]:
    FIG.mkdir(parents=True, exist_ok=True)
    _mpl()
    paths = {}

    stages = [
        "Fetched, latest view",
        "Curated subjects",
        "Accepted line diagrams",
        "Gold trajectories",
        "Drawn in Isaac Sim",
    ]
    values = [540, 186, 132, 39, 9]
    fig, ax = plt.subplots(figsize=(7.4, 3.5))
    y = list(range(len(stages)))[::-1]
    bars = ax.barh(y, values, color=BLUE_HEX, height=0.62)
    ax.set_yticks(y)
    ax.set_yticklabels(stages)
    for bar, value in zip(bars, values):
        ax.text(value + 8, bar.get_y() + bar.get_height() / 2, str(value),
                va="center", color=INK_HEX, fontsize=13)
    ax.set_xlim(0, 680)
    ax.set_xlabel("Samples")
    ax.xaxis.grid(True, color=GRID_HEX)
    ax.set_axisbelow(True)
    ax.tick_params(axis="y", length=0)
    ax.spines["left"].set_visible(False)
    fig.tight_layout()
    paths["funnel"] = FIG / "funnel.png"
    fig.savefig(paths["funnel"])
    plt.close(fig)

    metrics = [
        "Accepted by pipeline",
        "First-attempt pass",
        "Strong-referee pass",
        "BioCLIP top-1",
        "Exportable trajectory",
        "Finished within 10 min",
    ]
    flux = [65, 45, 100, 75, 100, 90]
    omni = [35, 5, 30, 30, 60, 55]
    fig, ax = plt.subplots(figsize=(8.6, 4.15))
    y = list(range(len(metrics)))[::-1]
    h = 0.36
    ax.barh([i + h / 2 for i in y], flux, height=h, color=BLUE_HEX, label="FLUX.1-schnell")
    ax.barh([i - h / 2 for i in y], omni, height=h, color=ORANGE_HEX, label="OmniGen2")
    ax.set_yticks(list(y))
    ax.set_yticklabels(metrics)
    ax.set_xlim(0, 118)
    ax.set_xticks([0, 25, 50, 75, 100])
    ax.set_xlabel("% of 20 frozen test subjects")
    ax.xaxis.grid(True, color=GRID_HEX)
    ax.set_axisbelow(True)
    ax.tick_params(axis="y", length=0)
    ax.spines["left"].set_visible(False)
    ax.legend(frameon=False, loc="lower left", bbox_to_anchor=(0.0, 1.02), ncol=2, fontsize=12, borderaxespad=0)
    fig.tight_layout()
    paths["r2"] = FIG / "r2_comparison.png"
    fig.savefig(paths["r2"])
    plt.close(fig)

    names = ["bird", "cat 70bc", "cat ab7b", "dog 96ab", "fish", "horse"]
    before = [1.554, 1.163, 1.495, 1.638, 1.604, 1.244]
    after = [0.315, 0.312, 0.314, 0.366, 0.394, 0.316]
    fig, ax = plt.subplots(figsize=(7.6, 3.7))
    y = list(range(len(names)))[::-1]
    h = 0.34
    ax.barh([i + h / 2 for i in y], before, height=h, color=ORANGE_HEX, label="Earlier controller")
    ax.barh([i - h / 2 for i in y], after, height=h, color=BLUE_HEX, label="baseline-v1")
    ax.set_yticks(list(y))
    ax.set_yticklabels(names)
    ax.set_xlim(0, 2.05)
    ax.set_xlabel("Tracking RMSE, desired path to executed path (mm)")
    ax.xaxis.grid(True, color=GRID_HEX)
    ax.set_axisbelow(True)
    ax.tick_params(axis="y", length=0)
    ax.spines["left"].set_visible(False)
    ax.legend(frameon=False, loc="lower left", bbox_to_anchor=(0.0, 1.02), ncol=2, fontsize=12, borderaxespad=0)
    fig.tight_layout()
    paths["controller"] = FIG / "controller.png"
    fig.savefig(paths["controller"])
    plt.close(fig)

    labels = ["Original", "Blur", "Resize\nhalf and back", "2% salt\nand pepper"]
    rates = [100, 0, 30, 0]
    colors = [BLUE_HEX, ORANGE_HEX, "#E0A15A", ORANGE_HEX]
    fig, ax = plt.subplots(figsize=(5.4, 3.15))
    bars = ax.bar(labels, rates, color=colors, width=0.68)
    for bar, rate in zip(bars, rates):
        ax.text(bar.get_x() + bar.get_width() / 2, rate + 3, f"{rate}%",
                ha="center", color=INK_HEX, fontsize=12)
    ax.set_ylim(0, 118)
    ax.set_ylabel("Prescreen pass rate")
    ax.yaxis.grid(True, color=GRID_HEX)
    ax.set_axisbelow(True)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0)
    fig.tight_layout()
    paths["robust"] = FIG / "validator_robustness.png"
    fig.savefig(paths["robust"])
    plt.close(fig)

    paths["dog"] = FIG / "dog_96ab6db6_strokes.png"
    _render_strokes(
        ROOT / "isaac-sim/inputs/svg/Dog/96ab6db6d9116afe.svg",
        paths["dog"],
        canvas=(689, 616),
    )
    return paths


def _render_strokes(svg_path: Path, out: Path, canvas: tuple[int, int]) -> None:
    text = svg_path.read_text(encoding="utf-8")
    polylines = re.findall(r'points="([^"]+)"', text)
    cw, ch = canvas
    scale = 900 / cw
    pad = 36
    img = Image.new("RGB", (int(cw * scale) + pad * 2, int(ch * scale) + pad * 2), "white")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([8, 8, img.width - 8, img.height - 8], radius=18, outline="#D7DEE7", width=2)
    for poly in polylines:
        pts = []
        for pair in poly.split():
            x, y = pair.split(",")
            pts.append((pad + float(x) * scale, pad + float(y) * scale))
        if len(pts) >= 2:
            draw.line(pts, fill="#1A1D21", width=3, joint="curve")
    img.save(out)


# ── slides ───────────────────────────────────────────────────────────────

def build() -> Path:
    figs = make_figures()
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "WildTrace Demo 1"
    prs.core_properties.subject = "Data preparation, analytics, and early model results"
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.55), Inches(0.18), Inches(2.15))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ORANGE
    accent.line.fill.background()
    _box(s, Inches(0.7), Inches(1.85), Inches(11), Inches(0.34), "DATA 298   ·   PROJECT DEMO 1", 15, True, RGBColor(0xF0, 0xC4, 0xA8))
    _box(s, Inches(0.7), Inches(2.25), Inches(11.5), Inches(0.9), "WildTrace", 60, True, WHITE)
    _box(s, Inches(0.7), Inches(3.25), Inches(11.2), Inches(1.1),
         "From a wildlife photograph to a pen trajectory\na robot arm can draw and we can measure.", 24, False, RGBColor(0xD6, 0xE2, 0xF0))
    _box(s, Inches(0.7), Inches(5.55), Inches(11), Inches(0.7),
         "Data preparation   ·   Train and test sets   ·   Analytics   ·   Early model results",
         16, False, RGBColor(0xA9, 0xC0, 0xD8))
    _notes(s, """
WildTrace turns wildlife photographs into short pen trajectories, then checks those trajectories on a simulated xArm7.
This demo covers the four things Demo 1 asks for: how the photos were prepared, how training and test data were separated, what the analytics show, and which models we have already run.
The numbers come from the project workbook, benchmark r2, and the frozen baseline-v1 robot runs. Where a model is implemented but not yet measured, we say so.
""")

    # 2 Problem
    s = prs.slides.add_slide(blank)
    _header(s, "The platform", "Two layers, one measured result")
    steps = [
        ("1", "Fetch", "Open Images photo\nand instance mask"),
        ("2", "Curate", "Quality, viewpoint,\nBioCLIP species"),
        ("3", "Draw", "FLUX outline,\nthen validate"),
        ("4", "Export", "At most 6 strokes,\ngold trajectory"),
        ("5", "Execute", "Isaac Sim xArm7,\ntracking error"),
    ]
    for i, (n, name, detail) in enumerate(steps):
        left = Inches(0.48 + i * 2.55)
        _rect(s, left, Inches(1.45), Inches(2.35), Inches(2.15), CARD, radius=0.1)
        _box(s, left + Inches(0.16), Inches(1.58), Inches(2.0), Inches(0.34), n, 14, True, ORANGE)
        _box(s, left + Inches(0.16), Inches(1.92), Inches(2.05), Inches(0.4), name, 20, True, NAVY)
        _box(s, left + Inches(0.16), Inches(2.4), Inches(2.05), Inches(0.9), detail, 14, False, MUTED)
    _rect(s, Inches(0.48), Inches(3.9), Inches(12.35), Inches(2.9), SOFT, radius=0.08)
    _box(s, Inches(0.75), Inches(4.1), Inches(11.8), Inches(0.4), "What this demo shows", 16, True, NAVY)
    _lines(s, Inches(0.75), Inches(4.6), Inches(11.8), Inches(2.0), [
        "Preprocessing: 540 fetched samples reduced to 186 drawable subjects, 132 accepted drawings, and 39 gold trajectories.",
        "Train and test: the curated corpus is the development set. Generator scores use a frozen 20-subject test slice, seed 14.",
        "Analytics: acceptance, pose fidelity, recognizability, judge agreement, and robot draw time on that same slice.",
        "Models already run: BioCLIP, FLUX.1-schnell, OmniGen2, three vision-language judges, and the baseline-v1 controller.",
    ], size=16, spacing=8)
    _footer(s, 2)
    _notes(s, """
The dataset layer prepares the drawing. The execution layer proves a robot can follow it.
A photo and its human-drawn mask come from Open Images. We keep only animals that are large enough, in focus, and seen from a drawable angle. BioCLIP names the species. FLUX redraws the animal as a clean outline. An OpenCV check and a vision-language judge accept or retry that drawing. The accepted outline becomes at most six pen strokes.
Isaac Sim then drives a UFACTORY xArm7 with a pen and records how far the pen tip strayed from the intended path.
The next slides walk through those results in the order the demo asks for.
""")

    # 3 Preprocessing method
    s = prs.slides.add_slide(blank)
    _header(s, "Data pre-processing", "What we keep, and what we drop")
    _rect(s, Inches(0.48), Inches(1.25), Inches(6.05), Inches(5.55), CARD, radius=0.08)
    _box(s, Inches(0.72), Inches(1.42), Inches(5.6), Inches(0.36), "Source", 18, True, NAVY)
    _lines(s, Inches(0.72), Inches(1.95), Inches(5.55), Inches(4.5), [
        "Open Images v7, images plus human instance masks, CC BY 2.0.",
        "One sample is one mask instance, not one file.",
        "Configured categories: Cat, Dog, Horse, Bird, Butterfly, and Fish at 100 each. Frog at 40.",
        "Butterfly has no train masks, so it contributes no samples.",
        "Fetch ledger: 646 rows. 641 downloaded. 5 failed on corrupt mask archives.",
        "Latest view used for curation: 540 samples.",
        "Official test-split URLs are configured and were not the reported corpus.",
    ], size=15, spacing=8)
    _rect(s, Inches(6.75), Inches(1.25), Inches(6.08), Inches(5.55), CARD, radius=0.08)
    _box(s, Inches(6.98), Inches(1.42), Inches(5.6), Inches(0.36), "Gates, in order", 18, True, NAVY)
    _lines(s, Inches(6.98), Inches(1.95), Inches(5.6), Inches(4.5), [
        "Drop duplicate images by checksum.",
        "Require at least 256 by 256 pixels. Resize to at most 1024.",
        "Drop blur below variance 0.001.",
        "Mask: coverage at least 0.03, border touch at most 0.50, at most 4 components, largest component at least 0.75.",
        "Viewpoint: keep profile, front, and front three-quarter. Drop rear, top-down, and occluded views.",
        "Viewpoint confidence at least 0.42 and margin at least 0.08.",
        "BioCLIP assigns the species label on the crop.",
    ], size=15, spacing=7)
    _footer(s, 3)
    _notes(s, """
Preprocessing starts from Open Images version 7. We use the photograph and the human segmentation mask. Each mask instance is one sample.
Seven animal categories were requested. Butterfly has no training masks in Open Images, so that category is empty. The fetch ledger recorded 646 attempts, 641 successful downloads, and 5 corrupt mask archives. After deduplication, the latest view holds 540 samples.
A sample survives only if the animal is large enough, sharp, mostly inside the frame, and seen from an angle a simple outline can express. Rear, top-down, and occluded views are rejected. BioCLIP then labels the species. The next slide is the count of what those gates kept.
""")

    # 4 Preprocessing results
    s = prs.slides.add_slide(blank)
    _header(s, "Data pre-processing", "540 samples become 39 robot-ready drawings")
    s.shapes.add_picture(str(figs["funnel"]), Inches(0.35), Inches(1.15), Inches(7.5), Inches(4.55))
    _kpi(s, Inches(8.05), Inches(1.25), Inches(4.75), Inches(1.15), "34.4%", "186 of 540 samples accepted as subjects")
    _kpi(s, Inches(8.05), Inches(2.55), Inches(4.75), Inches(1.15), "71.0%", "132 of 186 drawings passed validation")
    _kpi(s, Inches(8.05), Inches(3.85), Inches(4.75), Inches(1.15), "39", "Gold trajectories, 6 categories, 26 of 30 angle cells")
    _box(s, Inches(0.55), Inches(5.85), Inches(12.2), Inches(1.15),
         "Main rejection flags, and a sample can carry more than one: low mask coverage 169, top-heavy pose 141, blur 102.\n"
         "Of the 54 drawings that failed, every one used all 5 attempts and failed the OpenCV line check. Mean combined score of accepted drawings: 0.632.",
         14, False, MUTED)
    _footer(s, 4)
    _notes(s, """
This is the preprocessing result. Of 540 samples in the latest view, 186 were accepted as drawable subjects. That is 34.4 percent. The main rejection flags were low mask coverage, 169, a top-heavy pose, 141, and blur, 102. Those flags overlap, so they do not add up to the number rejected.
Those 186 subjects were redrawn as line diagrams. 132 were accepted, 71 percent. Fifty-four were rejected, and every rejected subject used all five attempts. They failed the OpenCV prescreen: too many components, too many small contours, or broken lower-body lines. The mean combined score of an accepted drawing is 0.632.
From the accepted drawings we keep the best one in each category, species, and angle bucket. That produces 39 gold trajectories across 6 categories, filling 26 of 30 angle cells. Nine distinct samples have been drawn on the simulated arm.
""")

    # 5 Samples
    s = prs.slides.add_slide(blank)
    _header(s, "Data pre-processing", "Accepted drawings the robot can be given")
    samples = [
        (ROOT / "isaac-sim/inputs/diagrams/Bird/c3a528aafbfd7b55.png", "Bird", "4 strokes"),
        (ROOT / "isaac-sim/inputs/diagrams/Cat/70bc30f8a5f918eb.png", "Cat", "5 strokes"),
        (ROOT / "isaac-sim/inputs/diagrams/Dog/ea59d6637578c29d.png", "Dog", "6 strokes"),
        (ROOT / "isaac-sim/inputs/diagrams/Fish/ba3e1657db195fe4.png", "Fish", "3 strokes"),
        (ROOT / "isaac-sim/inputs/diagrams/Frog/3f221bad695edcf8.png", "Frog", "2 strokes"),
        (ROOT / "isaac-sim/inputs/diagrams/Horse/7c54b4c75832b3b9.png", "Horse", "8 strokes*"),
    ]
    for i, (path, name, meta) in enumerate(samples):
        left = Inches(0.42 + i * 2.14)
        _rect(s, left, Inches(1.28), Inches(2.02), Inches(3.55), CARD, radius=0.08)
        s.shapes.add_picture(str(path), left + Inches(0.08), Inches(1.38), Inches(1.86), Inches(2.55))
        _box(s, left + Inches(0.08), Inches(4.0), Inches(1.86), Inches(0.32), name, 16, True, NAVY, align="center")
        _box(s, left + Inches(0.08), Inches(4.32), Inches(1.86), Inches(0.32), meta, 13, False, MUTED, align="center")
    _box(s, Inches(0.5), Inches(5.15), Inches(12.3), Inches(1.7),
         "Gold export keeps the longest contours, at most 6 strokes, and at most 128 points on each stroke. Coordinates are normalized to a 0–1 canvas, so the same file can drive the simulator without robot-specific units.\n"
         "Across 39 gold files: mean 4.46 strokes (range 2–6) and 379.5 points (range 155–768). Sixteen drawings sit on the 6-stroke cap.\n"
         "*The stored horse has 8 strokes from an earlier export. A new export of that drawing would fail the current 1-to-6 stroke gate.",
         14, False, INK)
    _footer(s, 5)
    _notes(s, """
These are accepted line drawings already used as robot inputs. Each one started as a photograph and a mask, passed curation and validation, and was converted to strokes.
The export rule is deliberate. We keep the longest contours, no more than six strokes, and no more than 128 points on a stroke. Points are normalized between zero and one, so the trajectory is not tied to one robot.
On the 39 gold files, the average drawing has 4.46 strokes and about 380 points. Sixteen of the 39 hit the six-stroke cap.
The horse on the right is a historical file with eight strokes. That file was drawn in simulation. Under the current export rule, a new export of the same drawing would be rejected. We keep the old file so the baseline run stays reproducible, and we do not treat eight strokes as the current contract.
""")

    # 6 Train / test
    s = prs.slides.add_slide(blank)
    _header(s, "Training and test data", "Development corpus, frozen test slice, robot baseline")
    _table(
        s, Inches(0.45), Inches(1.22), Inches(12.4), Inches(4.15),
        ["Set", "Size", "Role"],
        [
            ["Curated corpus", "540 → 186 subjects → 39 gold", "Development. Filters, generation, and export."],
            ["evalset_v2", "20 subjects, seed 14", "Paired generator test. Same images for both models."],
            ["evalset_v1", "30 subjects, 5 per category", "Per-model harness. Same seed."],
            ["labelset_v1", "60 diagrams: 44 accepted, 16 rejected", "Planned human labels. Not collected yet."],
            ["baseline-v1 draws", "8 gold trajectories, plus 1 later frog", "Frozen controller test and a reproduction check."],
            ["Behavior-cloning demos", "1 demonstration, 8,760 frames", "Training blocked until a second drawing exists."],
        ],
        [Inches(2.5), Inches(4.15), Inches(5.75)],
        size=13,
    )
    _box(s, Inches(0.5), Inches(5.55), Inches(12.3), Inches(1.4),
         "The 20-subject test is a frozen slice of the curated corpus, not a fresh download from the Open Images test split. Both generators see the identical subjects, so a difference is the model, not the photo.\n"
         "BioCLIP’s 0.885 mean confidence is not accuracy. There are no held-out species labels yet. The 235B referee is a proxy judge until labelset_v1 is labeled by hand.",
         14, False, MUTED)
    _footer(s, 6)
    _notes(s, """
Training and test data are separated by job, not by a single random split of every photo.
The development corpus is the 540-sample view, the 186 accepted subjects, and the 39 gold trajectories. That is what the filters and the export were built on.
Model comparison uses evalset_v2: 20 subjects, seed 14, three per category plus one extra cat and one extra dog. Both generators see those same 20 images. evalset_v1 is a 30-subject harness, five per category. labelset_v1 holds 60 diagrams, 44 the pipeline accepted and 16 it rejected, waiting for human labels. Those labels have not been collected, so we do not report a human accuracy number.
The robot baseline is eight gold trajectories drawn under one frozen configuration, plus a later frog used only as a reproduction check.
Behavior cloning has one demonstration, 8,760 frames of a two-stroke puppy. Training is blocked on purpose until a second drawing exists, so the policy cannot memorize a single path.
One caution: these evaluation subjects come from the curated corpus. They are frozen, and they are shared, but they are not the official Open Images test split.
""")

    # 7 Analytics corpus
    s = prs.slides.add_slide(blank)
    _header(s, "Data analytics", "What the processed corpus actually contains")
    tiles = [
        ("0.885", "Mean BioCLIP top-1\nconfidence on accepted subjects"),
        ("0.674", "Mean viewpoint\nconfidence"),
        ("51.1%", "Passed on the\nfirst generation attempt"),
        ("0.989", "R² of the draw-time\nmodel, 7 finished runs"),
    ]
    for i, (value, label) in enumerate(tiles):
        left = Inches(0.48 + (i % 4) * 3.2)
        _rect(s, left, Inches(1.28), Inches(3.02), Inches(1.7), CARD, radius=0.1)
        _box(s, left + Inches(0.16), Inches(1.4), Inches(2.7), Inches(0.5), value, 28, True, BLUE)
        _box(s, left + Inches(0.16), Inches(1.98), Inches(2.7), Inches(0.8), label, 14, False, MUTED)
    _rect(s, Inches(0.48), Inches(3.2), Inches(12.35), Inches(3.6), CARD, radius=0.08)
    _box(s, Inches(0.75), Inches(3.38), Inches(11.8), Inches(0.36), "How the 132 accepted drawings were earned", 18, True, NAVY)
    _lines(s, Inches(0.75), Inches(3.9), Inches(11.8), Inches(2.6), [
        "Attempts used by accepted drawings: 95 on try 1, 20 on try 2, 7 on try 3, 5 on try 4, 5 on try 5.",
        "Species confidence is uneven: Cat 0.771, Dog 0.810, Bird 0.904. That is model confidence, not a checked accuracy.",
        "OpenCV rejection counts on the 54 failures: too many components 49, too many small contours 47, lower-body fragments 34.",
        "Estimated draw time from seven finished robot runs: 0.3528 × pen-down millimeters + 26.28 × strokes + 4.52 seconds.",
        "First-attempt generation of all 186 subjects took 1.29 hours, median 24.2 seconds, on the laptop GPU with the model offloaded to CPU.",
    ], size=15, spacing=6)
    _footer(s, 7)
    _notes(s, """
These are the analytics on the data itself, before the head-to-head model test.
BioCLIP’s mean top-1 confidence on accepted subjects is 0.885. Viewpoint confidence averages 0.674. Confidence is not accuracy: we do not yet have human species labels to score against.
Just over half of the subjects, 51.1 percent, passed on the first drawing attempt. The rest of the 132 accepted drawings needed a retry. Ninety-five passed on attempt one, twenty on attempt two, and the remaining seventeen on attempts three through five.
The failures are structural. Of 54 rejects, 49 had too many components, 47 had too many small contours, and 34 had broken lines in the lower body.
We also fit a draw-time equation on seven finished simulator runs. Time in seconds is 0.3528 times pen-down millimeters, plus 26.28 times the stroke count, plus 4.52. R squared is 0.989, so stroke count and path length already predict how long the arm will take.
""")

    # 8 Generator analytics
    s = prs.slides.add_slide(blank)
    _header(s, "Data analytics", "Same 20 subjects, two generators")
    s.shapes.add_picture(str(figs["r2"]), Inches(0.25), Inches(1.12), Inches(8.15), Inches(5.55))
    _rect(s, Inches(8.4), Inches(1.35), Inches(4.5), Inches(5.15), CARD, radius=0.08)
    _box(s, Inches(8.6), Inches(1.52), Inches(4.15), Inches(0.7), "Shared referee\nQwen3-VL-235B", 16, True, NAVY)
    _lines(s, Inches(8.6), Inches(2.4), Inches(4.1), Inches(3.8), [
        "FLUX passed 20 of 20. OmniGen2 passed 6 of 20. Sign test p = 0.0001.",
        "BioCLIP recognized the animal on 75% of FLUX drawings and 30% of OmniGen2. p = 0.0004.",
        "Every FLUX drawing exported. 60% of OmniGen2 drawings did.",
        "Pose is the exception. OmniGen2’s non-blank sketches follow the photo more closely. Median silhouette IoU is about 0.2 for FLUX.",
    ], size=14, spacing=10)
    _footer(s, 8)
    _notes(s, """
Benchmark r2 ran both generators through the same retry loop, at most three attempts, on the same 20 subjects.
Read the strong referee row first. A 235-billion-parameter judge scored every final drawing after the run. FLUX passed 20 of 20. OmniGen2 passed 6 of 20. On the paired comparison FLUX won 18 subjects, OmniGen2 won 1, and 1 was a tie. The sign-test p-value is 0.0001.
An independent biology model, BioCLIP, named the right category for 75 percent of FLUX drawings and 30 percent of OmniGen2 drawings. p = 0.0004.
Every FLUX drawing became an exportable trajectory. Sixty percent of OmniGen2 drawings did.
The loop judges were not identical: FLUX was watched by the local 7B model, OmniGen2 by the hosted 8B model. That is why the shared 235B referee is the comparison we trust.
OmniGen2 is better at pose when it actually draws. Its silhouette overlap with the source mask is higher, and the difference is not significant at n = 20. FLUX tends to draw a generic stance. A lying cat can come back as a sitting cat. Median overlap for FLUX is about 0.2.
""")

    # 9 Failures
    s = prs.slides.add_slide(blank)
    _header(s, "Data analytics", "Three failure modes, and a check that the gate sees them")
    cards = [
        ("Blank page", "OmniGen2 left 9 of 20 final drawings empty. The local 7B judge called a blank “a simple cartoon” and passed it.", "OpenCV flag: low foreground ratio, under 1% dark pixels. The 235B referee rejected the blanks."),
        ("Changed pose", "FLUX passed the referee 20 of 20 and still replaced the photo’s pose. Facing-direction match was 65% for FLUX and 20% for OmniGen2.", "The prescreen does not score pose. Pose shows up in silhouette overlap and in the judge’s written reason."),
        ("Broken lines", "OmniGen2’s closest sketches, a betta at IoU 0.89 and a frog at 0.91, were rejected for too many pieces.", "Flags: too many components, small contours, lower-body fragments. A retry lowers strength so the next drawing stays nearer the contour."),
    ]
    for i, (title, body, detect) in enumerate(cards):
        top = Inches(1.18 + i * 1.42)
        _rect(s, Inches(0.42), top, Inches(8.15), Inches(1.32), CARD, radius=0.08)
        _box(s, Inches(0.6), top + Inches(0.08), Inches(7.8), Inches(0.3), title, 15, True, NAVY)
        _box(s, Inches(0.6), top + Inches(0.4), Inches(7.8), Inches(0.82), body + "  " + detect, 12, False, INK)
    s.shapes.add_picture(str(figs["robust"]), Inches(8.7), Inches(1.2), Inches(4.25), Inches(3.55))
    _box(s, Inches(8.75), Inches(4.8), Inches(4.15), Inches(1.9),
         "CPU check on 10 stored drawings. No new model was run. Blur and 2% speckle dropped the pass rate from 10/10 to 0/10. The same component gate that rejected fragmented sketches fired again.",
         13, False, MUTED)
    _footer(s, 9)
    _notes(s, """
Three failures show up again and again, and each one has a detector.
Blank pages: OmniGen2’s final drawing was empty for 9 of 20 subjects. The local 7-billion-parameter judge described a blank page as a simple cartoon and passed it. OpenCV catches this when dark pixels fall under 1 percent. The large referee also rejected every blank.
Changed pose: FLUX drawings look like clean animals, and the referee passed all 20, but they often do not match the photograph’s stance. Facing direction matched 65 percent of the time for FLUX and 20 percent for OmniGen2. The OpenCV gate does not measure pose. We measure it with silhouette overlap and direction scores.
Broken lines: OmniGen2 sometimes traced the animal very closely, including a betta at 0.89 overlap and a frog at 0.91, and was still rejected because the thin lines fell into many pieces.
We checked that sensitivity without calling a model. On 10 stored drawings, the original prescreen passed 10 of 10, mean score 0.57. Gaussian blur and 2 percent salt-and-pepper passed 0 of 10. Resizing to half and back passed 3 of 10. The gate is doing the job it was built for. It is not a pose checker.
""")

    # 10 ML selection
    s = prs.slides.add_slide(blank)
    _header(s, "Machine learning results", "What we ran, and what we kept")
    _table(
        s, Inches(0.4), Inches(1.2), Inches(12.5), Inches(4.55),
        ["Model", "Job", "Result", "Decision"],
        [
            ["BioCLIP", "Species label", "Mean confidence 0.885", "Keep for enrichment"],
            ["FLUX.1-schnell", "Outline generator", "Referee 20/20, export 20/20, 113 s", "Production default"],
            ["OmniGen2", "Outline generator", "Referee 6/20, 45% blank, better pose", "Not the default"],
            ["qwen2.5vl 7B", "Production judge", "Passed 40/40, including 9 blanks. κ = 0", "Replace"],
            ["qwen3-vl 8B", "Candidate judge", "Agrees with the 235B referee, κ = 0.83", "Use for the next run"],
            ["qwen3-vl 235B", "Referee", "Reference on all 40 finals", "Keep as the external check"],
            ["SDXL + ControlNet", "Outline generator", "Integrated. Quality not measured.", "Waiting on a GPU run"],
        ],
        [Inches(2.35), Inches(2.15), Inches(5.15), Inches(2.85)],
        size=13,
    )
    _box(s, Inches(0.48), Inches(5.95), Inches(12.3), Inches(0.95),
         "Cohen’s κ of 0 means the 7B judge matched the referee no better than chance, because it passed everything. The 8B judge agreed on 92.5% of 40 drawings. Switching judges does not require a new generator.",
         14, False, MUTED)
    _footer(s, 10)
    _notes(s, """
These are the model results we can stand behind.
BioCLIP labels species during curation. We keep it, and we do not quote its confidence as accuracy.
FLUX.1-schnell is the production outline generator. On the 20-subject test it passed the strong referee every time, exported every time, and took about 113 seconds per subject. OmniGen2 keeps the source pose more often when it draws, but 45 percent of its final drawings were blank, so it is not the default.
The judge result matters as much as the generator. The local 7B model passed all 40 final drawings, blanks included. Agreement with the 235B referee is a kappa of 0. In production, only the OpenCV prescreen was really rejecting anything. The hosted 8B judge agrees with that referee at kappa 0.83, 92.5 percent of drawings, and it rejected the blanks. The next comparison should keep FLUX and swap in the 8B judge.
SDXL with a scribble ControlNet is now in the codebase as a third generator, on the same validation and trajectory path. Its drawing quality has not been measured. Integration tests mock the network. We do not have a score to put in this table.
""")

    # 11 Robot
    s = prs.slides.add_slide(blank)
    _header(s, "Machine learning results", "The arm follows the drawing more closely")
    s.shapes.add_picture(str(figs["controller"]), Inches(0.2), Inches(1.15), Inches(7.7), Inches(5.15))
    _kpi(s, Inches(8.05), Inches(1.3), Inches(4.8), Inches(1.2), "0.330 mm", "Mean RMSE on 8 frozen baseline-v1 runs")
    _kpi(s, Inches(8.05), Inches(2.65), Inches(4.8), Inches(1.2), "4.3×", "On 6 paired runs: 1.45 mm down to 0.34 mm")
    _kpi(s, Inches(8.05), Inches(4.0), Inches(4.8), Inches(1.2), "8 / 8", "Finished. Longest draw, the horse, 563 seconds")
    _box(s, Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.7),
         "baseline-v1 is a frozen zero-gravity, contact-free measurement (config hash 15e6428dffca). It shows the controller is in the range of published pen-tracking errors. It is not a real-robot result. The current scene uses Earth gravity, 9.81 m/s².",
         13, False, MUTED)
    _footer(s, 11)
    _notes(s, """
The execution result is the frozen baseline-v1 controller. On six paired trajectories, tracking RMSE fell from about 1.45 millimeters to 0.34 millimeters, about 4.3 times lower. The chart is that pair. The full frozen set is eight runs, mean RMSE 0.330 millimeters, ranging from 0.270 to 0.394. Maximum error stayed between 0.97 and 1.17 millimeters. All eight finished. The longest, the eight-stroke horse, took 563 seconds, inside a 10-minute budget.
A later frog, drawn as a reproduction check, landed at 0.366 millimeters, inside that same range.
Two caveats belong with the number. These runs are zero-gravity and contact-free. Published pen-writing papers report roughly 0.4 to 1.5 millimeters on real hardware, so we are in a plausible range, and we are not claiming to beat those systems. The scene we train in now uses Earth gravity, 9.81 meters per second squared. A zero-gravity checkpoint must not be resumed under gravity, because the dynamics changed.
""")

    # 12 Ongoing
    s = prs.slides.add_slide(blank)
    _header(s, "Ongoing learning", "What is training, and what is only integrated")
    cols = [
        ("PPO on the arm", [
            "Two 300-update runs finished, about 19.5 minutes each.",
            "Step reward sat around 0.08 to 0.25. Tracking near the target was often 1.4 to 5.4 mm.",
            "No clear improvement. The episode had 800 steps and the path had 957 targets, so finishing was impossible.",
            "The limit is now 3,000 steps, with a four-stage curriculum from 100 targets at 3 mm down to the full path at 2 mm.",
        ]),
        ("Behavior cloning", [
            "Network is ready: layers 256, 256, 128, 36 inputs, a Cartesian step of at most 5 mm.",
            "Loss is mean squared error, batch 128, early stopping.",
            "One demonstration exists. Training stays off until a second drawing ID is recorded.",
            "The 39 gold trajectories are the planned demonstration source.",
        ]),
        ("Platform right now", [
            "Bronze, silver, and gold stages run as separate scripts with lineage and a config hash.",
            "SDXL + scribble ControlNet is registered beside FLUX. Inference has not been run.",
            "A shared contour bug was fixed: masks were thresholded at 128 and came out blank. The cut is now any positive pixel.",
            "Benchmark r2 is the record from before that fix. It is not a retest of the corrected code.",
        ]),
    ]
    for i, (title, bullets) in enumerate(cols):
        left = Inches(0.4 + i * 4.28)
        _rect(s, left, Inches(1.22), Inches(4.1), Inches(5.55), CARD, radius=0.08)
        _box(s, left + Inches(0.18), Inches(1.38), Inches(3.75), Inches(0.4), title, 16, True, NAVY)
        _lines(s, left + Inches(0.18), Inches(1.9), Inches(3.75), Inches(4.6), bullets, size=13, spacing=8)
    _footer(s, 12)
    _notes(s, """
Three things are in progress, and only one of them has a training log.
Proximal policy optimization has completed two sessions of 300 updates, about nineteen and a half minutes each. Reward per step was roughly 0.08 to 0.25, and the pen was often 1.4 to 5.4 millimeters from the active target. We could not show that the policy was improving. The reason was in the setup: the path had 957 targets and the episode allowed 800 steps, so the arm could never finish. Episodes are now 3,000 steps, and training advances through a curriculum, from 100 targets at a 3-millimeter tolerance up to the full path at 2 millimeters. The best checkpoint is kept separate from the latest one, and it will be compared with inverse kinematics alone.
Behavior cloning is implemented and not trained. The network is a three-layer perceptron. We have one demonstration. Training is blocked until there are two, so we do not report a cloning accuracy.
On the platform, the medallion pipeline is in place. SDXL with ControlNet is wired into the same validation loop and has not generated a scored drawing. While testing it, we found that the shared contour helper treated a 0-1 mask as empty. That bug is fixed. Because the fix changes conditioning for FLUX as well, the r2 numbers stay the historical record. They are not a score for the corrected code.
""")

    # 13 Walkthrough
    s = prs.slides.add_slide(blank)
    _header(s, "One sample", "Dog 96ab6db6, from strokes to a measured draw")
    _rect(s, Inches(0.42), Inches(1.2), Inches(6.35), Inches(5.55), CARD, radius=0.08)
    s.shapes.add_picture(str(figs["dog"]), Inches(0.7), Inches(1.4), Inches(5.8), Inches(4.55))
    _box(s, Inches(0.7), Inches(6.05), Inches(5.8), Inches(0.5),
         "Exported strokes, not the source photo. 4 strokes, 210 points.", 13, False, MUTED, align="center")
    facts = [
        ("Identity", "Dog, sample 96ab6db6. Pipeline status accepted, attempt 01."),
        ("Trajectory", "Canvas 689 × 616. Normalized 0–1. Inside the 6-stroke cap."),
        ("baseline-v1", "RMSE 0.366 mm. Max error 1.084 mm. Draw time 380 seconds."),
        ("Separate run", "The two-stroke puppy at Earth gravity is a different target: pen-down RMSE 0.738 mm. Do not merge it with this dog."),
    ]
    for i, (label, text) in enumerate(facts):
        top = Inches(1.25 + i * 1.3)
        _box(s, Inches(7.05), top, Inches(5.7), Inches(0.3), label, 14, True, ORANGE)
        _box(s, Inches(7.05), top + Inches(0.32), Inches(5.7), Inches(0.85), text, 15, False, INK)
    _footer(s, 13)
    _notes(s, """
This is one subject pulled all the way through.
Dog 96ab6db6 was accepted on the first drawing attempt. The figure is the exported pen path: four strokes, 210 points, on a 689 by 616 canvas, normalized between zero and one. Four strokes satisfies the current export cap of six.
Under baseline-v1 the simulated arm drew that trajectory with a tracking RMSE of 0.366 millimeters, a max error of 1.084 millimeters, and a draw time of 380 seconds. The config hash on that family of runs is 15e6428dffca, so the number belongs with the frozen baseline, not with a later code change.
A different file, the two-stroke puppy, has been run in the current Earth-gravity scene. Its pen-down RMSE is about 0.74 millimeters. That is a status check on the present simulator. It is not a replay of this dog, and the two errors should stay apart.
Every derived record carries the bronze sample id, the input hashes, and the config hash. Robot runs add the git SHA. That is how a slide number can be traced back to a file.
""")

    # 14 Close
    s = prs.slides.add_slide(blank)
    _header(s, "Demo 1", "Four results, and the honest gap")
    closes = [
        ("Prepared", "540 → 186 → 132 → 39", "Photos filtered by mask, focus, and viewpoint, then exported as normalized strokes."),
        ("Split", "20-subject frozen test", "Both generators scored on evalset_v2. Human labels and the official test split are still ahead."),
        ("Analyzed", "Referee 20/20 vs 6/20", "FLUX is the recognizable, exportable generator. Pose match and blank pages are the open defects."),
        ("Learned", "0.330 mm mean RMSE", "Controller baseline is measured. PPO and behavior cloning are not ready to claim a gain."),
    ]
    for i, (kicker, value, body) in enumerate(closes):
        left = Inches(0.42 + (i % 2) * 6.45)
        top = Inches(1.25 + (i // 2) * 2.15)
        _rect(s, left, top, Inches(6.2), Inches(2.0), CARD, radius=0.1)
        _box(s, left + Inches(0.22), top + Inches(0.14), Inches(5.7), Inches(0.28), kicker.upper(), 12, True, ORANGE)
        _box(s, left + Inches(0.22), top + Inches(0.46), Inches(5.7), Inches(0.45), value, 22, True, NAVY)
        _box(s, left + Inches(0.22), top + Inches(1.05), Inches(5.7), Inches(0.75), body, 14, False, INK)
    _box(s, Inches(0.5), Inches(5.7), Inches(12.3), Inches(1.2),
         "Next measured step: rerun FLUX and OmniGen2 after the contour fix, with the 8B judge, and add SDXL on the same 20 subjects. Then compare the 1 g policy with inverse kinematics alone.",
         15, False, MUTED)
    _footer(s, 14)
    _notes(s, """
To close on the four Demo 1 questions.
Data preparation is done through gold export. Five hundred forty samples became 186 subjects, 132 accepted drawings, and 39 trajectories.
Training and test data are prepared as a development corpus plus frozen evaluation sets. The generator test is 20 shared subjects. We still need human labels, and we have not scored the official Open Images test split.
The analytics say FLUX is the generator to keep for recognizability and for trajectories the robot can draw, and that the current 7B judge should be replaced. Pose drift and OmniGen2’s blank pages are measured defects, not anecdotes.
The machine-learning result we can report today is the controller baseline, mean tracking error 0.33 millimeters on eight runs, and the generator and judge comparison. PPO has been trained, and it has not yet shown a gain. Behavior cloning has not been trained. SDXL is integrated and unscored.
The next run that would move the project is a three-generator comparison on the same 20 subjects, after the contour fix, with one shared judge.
""")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
